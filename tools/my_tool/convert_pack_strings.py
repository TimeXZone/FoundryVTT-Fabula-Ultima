import argparse
import csv
from pathlib import Path
from openpyxl import load_workbook

ENTRIES_SHEET = "Entries"
TRANSLATIONS_SHEET = "Translations"


def sheet_to_csv(ws, out_path: Path):
    with out_path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        for row in ws.iter_rows(values_only=True):
            writer.writerow(["" if c is None else c for c in row])


def main():
    parser = argparse.ArgumentParser(description="Convert pack_strings.xlsx into two CSV files")
    parser.add_argument("--input", type=Path, default=Path("pack_strings.xlsx"), help="Input XLSX file")
    parser.add_argument("--out-dir", type=Path, default=Path(".") / "i18n", help="Output directory for CSVs")
    args = parser.parse_args()

    wb_path: Path = args.input
    out_dir: Path = args.out_dir
    if not wb_path.exists():
        raise SystemExit(f"Input workbook not found: {wb_path}")
    out_dir.mkdir(parents=True, exist_ok=True)

    wb = load_workbook(wb_path, read_only=True)

    # Entries sheet -> entries.csv
    if ENTRIES_SHEET in wb.sheetnames:
        ws = wb[ENTRIES_SHEET]
        out_path = out_dir / "entries.csv"
        sheet_to_csv(ws, out_path)
        print(f"Wrote {out_path}")
    else:
        print(f"Warning: sheet '{ENTRIES_SHEET}' not found in workbook")

    # Translations sheet -> translations.csv
    if TRANSLATIONS_SHEET in wb.sheetnames:
        ws = wb[TRANSLATIONS_SHEET]
        out_path = out_dir / "translations.csv"
        sheet_to_csv(ws, out_path)
        print(f"Wrote {out_path}")
    else:
        print(f"Warning: sheet '{TRANSLATIONS_SHEET}' not found in workbook")


if __name__ == "__main__":
    main()
