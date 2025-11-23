import argparse
import json
from pathlib import Path
from typing import Dict, List

from openpyxl import load_workbook

ENTRIES_SHEET = "Entries"
TRANSLATIONS_SHEET = "Translations"


def parse_identifier(identifier: str):
    try:
        file_part, pointer = identifier.split("::", 1)
    except ValueError:
        raise ValueError(f"Invalid identifier format: {identifier}")
    path_segments = [segment for segment in pointer.strip("/").split("/") if segment]
    return file_part, path_segments


def set_value(node, path: List[str], value):
    if not path:
        return value
    head, *tail = path
    key = int(head) if head.isdigit() else head
    if isinstance(key, int):
        if not isinstance(node, list) or key >= len(node):
            raise KeyError(f"List index out of range for path segment '{head}'")
        node[key] = set_value(node[key], tail, value)
    else:
        if not isinstance(node, dict) or key not in node:
            raise KeyError(f"Missing key '{key}' in object while setting value")
        node[key] = set_value(node[key], tail, value)
    return node


def load_entries_sheet(workbook_path: Path):
    wb = load_workbook(workbook_path)
    if ENTRIES_SHEET not in wb.sheetnames or TRANSLATIONS_SHEET not in wb.sheetnames:
        raise ValueError(
            f"Workbook must contain sheets '{ENTRIES_SHEET}' and '{TRANSLATIONS_SHEET}'"
        )
    entries_sheet = wb[ENTRIES_SHEET]
    translations_sheet = wb[TRANSLATIONS_SHEET]
    entries = []
    for row in entries_sheet.iter_rows(min_row=2, values_only=True):
        identifier, english = row[:2]
        if identifier and english:
            entries.append((identifier, english))
    translations: Dict[str, str] = {}
    for row in translations_sheet.iter_rows(min_row=2, values_only=True):
        english, chinese = row[:2]
        if english and chinese:
            translations[english] = chinese
    return entries, translations


def apply_translations(base_dir: Path, workbook_path: Path, output_dir: Path):
    entries, translations = load_entries_sheet(workbook_path)
    files_cache: Dict[Path, Dict] = {}
    modifications: Dict[Path, bool] = {}

    for identifier, english in entries:
        if english not in translations:
            continue
        translation = translations[english]
        file_part, path_segments = parse_identifier(identifier)
        json_path = base_dir / file_part
        if json_path not in files_cache:
            with json_path.open("r", encoding="utf-8") as f:
                files_cache[json_path] = json.load(f)
            modifications[json_path] = False
        try:
            set_value(files_cache[json_path], path_segments, translation)
            modifications[json_path] = True
        except KeyError as exc:
            print(f"Skipping {identifier}: {exc}")
            continue

    for json_path, modified in modifications.items():
        if not modified:
            continue
        target_path = json_path
        if output_dir:
            target_path = output_dir / json_path.relative_to(base_dir)
            target_path.parent.mkdir(parents=True, exist_ok=True)
        with target_path.open("w", encoding="utf-8") as f:
            json.dump(files_cache[json_path], f, ensure_ascii=False, indent=2)
            f.write("\n")
        print(f"Updated {target_path}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description=(
            "Apply translations from an XLSX workbook back into JSON files. "
            "Requires sheets named 'Entries' and 'Translations'."
        )
    )
    parser.add_argument(
        "--input",
        type=Path,
        default=Path("src/packs"),
        help="Base directory containing original JSON files (default: src/packs)",
    )
    parser.add_argument(
        "--workbook",
        type=Path,
        default=Path("pack_strings.xlsx"),
        help="Workbook containing extracted entries and translations (default: pack_strings.xlsx)",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=None,
        help=(
            "Optional output directory for translated JSON. If omitted, files are overwritten "
            "in place under the input directory."
        ),
    )
    args = parser.parse_args()
    apply_translations(args.input, args.workbook, args.output)